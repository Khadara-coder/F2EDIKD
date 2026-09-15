"""Generate the Excel inventory of File2EDI validation rules."""

from collections import Counter
from datetime import datetime
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(os.environ.get(
    "VALIDATION_RULES_XLSX_OUTPUT",
    str(ROOT / "docs" / "validation_rules_inventory.xlsx"),
))
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.rejection_catalog import (
    CODE_ALIASES,
    ISSUE_TAXONOMY,
    REJECTION_CATALOG,
    action_text,
    format_rejection_message,
    issue_taxonomy,
    review_actions,
)

HEADERS = [
    "Code canonique", "Alias / anciens codes", "Nom de la règle", "Description",
    "Domaine", "Étape du pipeline", "Nature du contrôle", "Portée",
    "Bloquante ?", "Sévérité", "Revue utilisateur requise ?", "Retry autorisé ?",
    "Statut métier", "Action automatique", "Action corrective", "Message français",
    "Message anglais", "Fichier / module", "Fonction", "Tests associés",
    "Existe dans le code ?", "Utilisée actuellement ?", "Observations",
]

HEADERS_AFFICHAGES = [
    "Ordre d'affichage", "Groupe / Domaine", "Libellé Groupe UI", "Code Anomalie",
    "Message affiché (FR)", "Badge Sévérité", "Bloquante ?",
    "Bouton Accepter / Valider", "Action au clic Accepter",
    "Bouton Refuser / Rejeter", "Action au clic Refuser",
    "Mode d'action", "Revue requise ?",
    "Zone de correction dans l'écran Revue", "Conseil d'action opérateur",
]

DISPLAY_DOMAIN_ORDER = [
    ("DOCUMENT", "Document", 1),
    ("PARTNER", "Partenaire", 2),
    ("DUPLICATE", "Doublon", 3),
    ("ARTICLE", "Article", 4),
    ("ORDER", "Commande", 5),
    ("EDI", "EDI", 6),
    ("DELIVERY", "Livraison", 7),
    ("TECHNICAL", "Technique", 8),
]

DOMAIN_ORDER_MAP = {d[0]: (d[2], d[1]) for d in DISPLAY_DOMAIN_ORDER}

SEVERITY_DISPLAY = {
    "CRITICAL": "Critique",
    "ERROR": "Erreur",
    "WARNING": "Avertissement",
    "INFO": "Info",
}

SEVERITY_SORT = {
    "CRITICAL": 1,
    "ERROR": 2,
    "WARNING": 3,
    "INFO": 4,
}

UI_ZONES = {
    "DOCUMENT": "Panneau Aperçu PDF / Action Retraiter",
    "PARTNER": "Informations générales (Sold-to / Ship-to)",
    "DUPLICATE": "Tableau des anomalies / En-tête commande",
    "ARTICLE": "Tableau Lignes de commande",
    "ORDER": "Informations générales (N° commande, Dates)",
    "EDI": "Validation & Génération EDIFACT",
    "DELIVERY": "Envoi vers SAP / Livraison SFTP",
    "TECHNICAL": "Données Maîtres / Système",
}

ENGINE_CODES = {
    "NO_DELIVERY_ADDRESS", "SHIPTO_NO_STRONG_MATCH", "ARTICLE_NOT_FOUND",
    "QUANTITY_MISSING", "UNIT_PRICE_MISSING", "ORDER_KEY_MISSING",
    "PO_NUMBER_DUPLICATE", "SOLDTO_NOT_FOUND", "NOT_AN_ORDER",
    "NO_LINE_ITEMS", "ORDER_CHANGE",
}

FUNCTIONS = {
    "SHIPTO_NO_STRONG_MATCH": "_check_delivery_address",
    "NO_DELIVERY_ADDRESS": "_check_delivery_address",
    "ORDER_KEY_MISSING": "_check_po_number",
    "PO_NUMBER_DUPLICATE": "_check_po_duplicate",
    "SOLDTO_NOT_FOUND": "_check_customer",
    "QUANTITY_MISSING": "_check_line_items",
    "UNIT_PRICE_MISSING": "_check_line_items",
    "ARTICLE_NOT_FOUND": "_check_line_items",
    "NO_LINE_ITEMS": "_check_line_items",
    "NOT_AN_ORDER": "_check_document_type",
    "ORDER_CHANGE": "_check_document_type",
}

NATURES = {
    "DOCUMENT": "Structurelle / classification",
    "PARTNER": "Référentiel / matching",
    "ORDER": "Complétude / syntaxe",
    "ARTICLE": "Référentiel / métier",
    "EDI": "Structurelle / cohérence EDI",
    "DUPLICATE": "Métier / idempotence",
    "DELIVERY": "Intégration",
    "TECHNICAL": "Technique",
}

MODULES = {
    "DOCUMENT": "app/extraction.py",
    "PARTNER": "app/engines/rejection_engine.py",
    "ORDER": "app/extraction.py",
    "ARTICLE": "app/engines/rejection_engine.py; src/pompac_rules.py",
    "EDI": "app/edifact_generator.py",
    "DUPLICATE": "src/file2edi/store.py",
    "DELIVERY": "src/sftp_delivery.py",
    "TECHNICAL": "src/rejection_catalog.py",
}

SPECIAL_ROWS = [
    {
        "code": "SOLDTO_CONFIDENCE_MIN", "name": "Seuil de confiance Sold-to",
        "description": "Critère documenté de matching ; aucun code applicatif canonique trouvé.",
        "domain": "PARTNER", "stage": "MATCHING", "nature": "Matching",
        "scope": "ORDER", "blocking": "À vérifier", "severity": "WARNING",
        "review": "Oui", "retry": "Oui", "status": "PENDING_USER_INPUT",
        "auto": "Mettre en revue si implémenté", "action": "Confirmer le client",
        "fr": "Confiance Sold-to insuffisante", "en": "Sold-to confidence below threshold",
        "module": "Documentation uniquement", "function": "", "used": "Non",
        "observations": "Critère documenté, pas un code du catalogue.",
    },
    {
        "code": "SHIPTO_CONFIDENCE_MIN", "name": "Seuil de confiance Ship-to",
        "description": "Critère documenté de matching ; aucun code applicatif canonique trouvé.",
        "domain": "PARTNER", "stage": "MATCHING", "nature": "Matching",
        "scope": "DELIVERY", "blocking": "À vérifier", "severity": "WARNING",
        "review": "Oui", "retry": "Oui", "status": "PENDING_USER_INPUT",
        "auto": "Mettre en revue si implémenté", "action": "Confirmer l'adresse",
        "fr": "Confiance Ship-to insuffisante", "en": "Ship-to confidence below threshold",
        "module": "Documentation uniquement", "function": "", "used": "Non",
        "observations": "Critère documenté, pas un code du catalogue.",
    },
]


def aliases_by_code():
    result = {code: [] for code in REJECTION_CATALOG}
    for alias, canonical in CODE_ALIASES.items():
        result.setdefault(canonical, []).append(alias)
    return result


def action_for(code, entry, taxonomy):
    if code in {"DELIVERY_SFTP_FAILED", "DELIVERY_EMAIL_FAILED", "PDF_PARSE_FAILURE"}:
        return "Relancer automatiquement puis escalader"
    if taxonomy["blocking"]:
        return "Bloquer le traitement"
    if entry["manual_review_required"]:
        return "Mettre en revue"
    return "Journaliser et continuer"


def rule_rows():
    aliases = aliases_by_code()
    rows = []
    for code, entry in REJECTION_CATALOG.items():
        taxonomy = ISSUE_TAXONOMY[code]
        rows.append({
            "code": code,
            "aliases": ", ".join(aliases.get(code, [])),
            "name": code.replace("_", " ").title(),
            "description": entry["message_fr"],
            "domain": taxonomy["domain"], "stage": taxonomy["stage"],
            "nature": NATURES[taxonomy["domain"]], "scope": taxonomy["scope"],
            "blocking": "Oui" if taxonomy["blocking"] else "Non",
            "severity": taxonomy["issue_severity"],
            "review": "Oui" if taxonomy["requires_user_input"] else "Non",
            "retry": "Oui" if entry["retry_allowed"] else "Non",
            "status": entry["business_status"],
            "auto": action_for(code, entry, taxonomy),
            "action": "Corriger les données, le document ou le référentiel selon le contexte",
            "fr": entry["message_fr"], "en": entry["message_en"],
            "module": MODULES[taxonomy["domain"]],
            "function": FUNCTIONS.get(code, ""),
            "tests": "tests/test_rejection_catalog.py" if code in REJECTION_CATALOG else "",
            "exists": "Oui", "used": "Oui" if code in ENGINE_CODES else "À vérifier",
            "observations": "Alias normalisés : " + ", ".join(aliases[code]) if aliases.get(code) else "",
        })
    return rows + SPECIAL_ROWS


def display_rows():
    rows = []
    for code, entry in REJECTION_CATALOG.items():
        tax = issue_taxonomy(code)
        act = review_actions(code)
        txt = action_text(code, lang="fr")
        domain = tax["domain"]
        order_num, domain_label = DOMAIN_ORDER_MAP.get(domain, (99, domain))

        rows.append({
            "order_num": order_num,
            "order_display": f"{order_num}. {domain_label}",
            "domain": domain,
            "domain_label": domain_label,
            "code": code,
            "message_fr": entry["message_fr"],
            "severity_display": SEVERITY_DISPLAY.get(tax["issue_severity"], tax["issue_severity"]),
            "severity_sort": SEVERITY_SORT.get(tax["issue_severity"], 99),
            "blocking": "Oui" if tax["blocking"] else "Non",
            "button_accept": act["button_accept"],
            "auto_action_accept": act["auto_action_accept"],
            "button_reject": act["button_reject"],
            "auto_action_reject": act["auto_action_reject"],
            "mode": act["mode"],
            "review": "Oui" if tax["requires_user_input"] else "Non",
            "ui_zone": UI_ZONES.get(domain, "Écran Revue"),
            "action_text": txt,
        })

    rows.sort(key=lambda r: (r["order_num"], r["severity_sort"], r["code"]))
    return rows


def write_table(ws, title, headers, rows, table_name, widths=None, blocking_col="I"):
    ws.append([title])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    end_row = ws.max_row
    end_col = ws.max_column
    ref = f"A2:{chr(64 + end_col)}{end_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ref
    ws.row_dimensions[1].height = 26
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + index)].width = width
    if blocking_col:
        ws.conditional_formatting.add(
            f"{blocking_col}3:{blocking_col}{end_row}",
            CellIsRule(operator="equal", formula=['"Oui"'], fill=PatternFill("solid", fgColor="F4CCCC")),
        )
    return end_row


def main():
    rows = rule_rows()
    disp_rows = display_rows()
    workbook = Workbook()

    # 1. Onglet Règles (Catalogue technique & métier complet)
    rules = workbook.active
    rules.title = "Règles"
    values = [[row.get(key, "") for key in (
        "code", "aliases", "name", "description", "domain", "stage", "nature", "scope",
        "blocking", "severity", "review", "retry", "status", "auto", "action", "fr",
        "en", "module", "function", "tests", "exists", "used", "observations",
    )] for row in rows]
    rules_widths = [28, 24, 28, 48, 16, 24, 28, 14, 12, 14, 20, 16, 22, 30, 42, 48, 48, 34, 30, 34, 18, 20, 44]
    write_table(rules, "Inventaire des règles de validation File2EDI", HEADERS, values, "ValidationRules", widths=rules_widths, blocking_col="I")

    # 2. Onglet Affichages (Anomalies affichées lors du traitement, boutons et actions)
    affichages = workbook.create_sheet("Affichages")
    affichages_values = [[row.get(key, "") for key in (
        "order_display", "domain", "domain_label", "code", "message_fr",
        "severity_display", "blocking", "button_accept", "auto_action_accept",
        "button_reject", "auto_action_reject", "mode", "review",
        "ui_zone", "action_text",
    )] for row in disp_rows]
    affichages_widths = [18, 16, 18, 28, 55, 16, 12, 30, 35, 30, 35, 15, 15, 38, 55]
    write_table(affichages, "Anomalies affichées lors du traitement, boutons et actions de revue", HEADERS_AFFICHAGES, affichages_values, "AnomaliesAffichage", widths=affichages_widths, blocking_col="G")

    # 3. Onglet Synthèse
    summary = workbook.create_sheet("Synthèse")
    summary.append(["Synthèse du catalogue de validations"])
    summary.append(["Métrique", "Valeur"])
    summary_data = [
        ("Nombre total de règles inventoriées", len(rows)),
        ("Règles canoniques du catalogue", len(REJECTION_CATALOG)),
        ("Critères documentés non canoniques", len(SPECIAL_ROWS)),
        ("Règles bloquantes", sum(row["blocking"] == "Oui" for row in rows)),
        ("Règles non bloquantes", sum(row["blocking"] == "Non" for row in rows)),
        ("Règles à vérifier", sum(row["blocking"] == "À vérifier" for row in rows)),
    ]
    for item in summary_data:
        summary.append(list(item))
    summary.append([])
    summary.append(["Répartition par domaine (ordre d'affichage UI)", "Nombre"])
    for domain_code, domain_label, order_idx in DISPLAY_DOMAIN_ORDER:
        count = sum(1 for r in disp_rows if r["domain"] == domain_code)
        summary.append([f"{order_idx}. {domain_label} ({domain_code})", count])
    summary.append([])
    summary.append(["Répartition par sévérité", "Nombre"])
    for key, count in sorted(Counter(row["severity"] for row in rows).items()):
        summary.append([key, count])
    summary.freeze_panes = "A3"
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 18
    summary["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    summary.merge_cells("A1:B1")
    for cell in summary[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")

    # 4. Onglet Incohérences
    issues = workbook.create_sheet("Incohérences")
    issue_headers = ["Sujet", "Constat", "Statut", "Action recommandée"]
    issue_rows = [
        ["ARTICLE_NOT_FOUND", "Le moteur renvoie maintenant blocking, conforme à la taxonomie.", "Corrigé", "Conserver le test de régression."],
        ["NO_LINE_ITEMS", "Le moteur renvoie maintenant blocking, car aucun EDIFACT ne peut être généré sans ligne.", "Corrigé", "Conserver le test de régression."],
        ["QUANTITY_INVALID", "Alias normalisé vers ARTICLE_QUANTITY_INVALID.", "Corrigé", "Utiliser le code canonique dans les nouveaux écrans et rapports."],
        ["PRICE_MISSING", "Alias normalisé vers UNIT_PRICE_MISSING.", "Corrigé", "Utiliser UNIT_PRICE_MISSING partout."],
        ["Codes PARTNER consolidés", "7 codes canoniques clairs avec gestion de famille et mismatch.", "Corrigé", "Vérification stricte de l'appartenance Ship-to/Sold-to."],
        ["SOLDTO_CONFIDENCE_MIN / SHIPTO_CONFIDENCE_MIN", "Critères présents dans la documentation mais absents du code et du catalogue.", "À clarifier", "Les traiter comme seuils de matching ou les implémenter explicitement."],
        ["Codes catalogue non détectés dans rejection_engine", "Certains codes sont gérés par d'autres modules ou restent à vérifier.", "À vérifier", "Ajouter un test d'intégration par code lorsque le flux sera stabilisé."],
    ]
    write_table(issues, "Points de cohérence et actions", issue_headers, issue_rows, "ValidationIssues", widths=[30, 50, 16, 45], blocking_col=None)

    thin = Side(style="thin", color="D9E2F3")
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(bottom=thin)
        ws.sheet_view.showGridLines = False

    workbook.properties.title = "Inventaire des règles de validation File2EDI"
    workbook.properties.creator = "File2EDI"
    workbook.properties.created = datetime.now()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT
    try:
        workbook.save(output_path)
    except PermissionError:
        output_path = OUTPUT.with_name(f"{OUTPUT.stem}_updated{OUTPUT.suffix}")
        workbook.save(output_path)
        print(f"Avertissement: {OUTPUT.name} est verrouillé; fichier écrit dans {output_path.name}")
    print(output_path)
    print(f"{len(rows)} lignes de règles générées")


if __name__ == "__main__":
    main()